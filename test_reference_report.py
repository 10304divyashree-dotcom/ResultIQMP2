"""Test report generation with the reference .xlsx file."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from services.analyser import analyse_workbook
from services.reporter import generate_report
import openpyxl

filepath = 'Reference/Original_1 ExcelSpreadsheet.xlsx'
output = 'reports/TEST_FIX_reference_report.xlsx'

# Step 1: Analyse
print("=== Analysing workbook ===")
worksheets = analyse_workbook(filepath)
for s in worksheets:
    sheet_name = s.get('name', '?')
    hdr = s.get('header_row_index', '?')
    marks = [c['name'] for c in s.get('columns', []) if c.get('classification') == 'marks']
    print(f"  Sheet: {sheet_name}, header_row={hdr}, marks_cols={marks}")

# Step 2: Generate report
print("\n=== Generating report ===")
sheet_name = worksheets[0]['name']
try:
    stats = generate_report(filepath, worksheets, sheet_name, output)
    print(f"  Success! Stats: total={stats['total_students']}, "
          f"pass={stats['pass_count']}, fail={stats['fail_count']}, "
          f"pass_rate={stats['pass_rate']}%")
except Exception as e:
    print(f"  ERROR: {e}")
    import traceback
    traceback.print_exc()
    sys.exit(1)

# Step 3: Inspect
print("\n=== Inspecting output ===")
wb = openpyxl.load_workbook(output, data_only=False)
ws = wb[sheet_name]
print(f"  max_row={ws.max_row}, max_col={ws.max_column}")
for r in range(max(1, ws.max_row - 25), ws.max_row + 1):
    row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    non_empty = [v for v in row_vals if v is not None and str(v).strip() != '']
    if non_empty:
        print(f"  Row {r}: {row_vals}")

print("\nDone!")
