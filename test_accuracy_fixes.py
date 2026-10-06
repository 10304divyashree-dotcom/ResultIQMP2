"""
Test accuracy improvements on sample files and edge cases.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
import openpyxl

# Test 1: Column Classification with new rules
from services.analyser import classify_column

print("--- Testing Classifications ---")
test_cases = [
    (pd.Series([101, 102, 103, 104, 105]), 'Roll', 'identity'),
    (pd.Series([1, 2, 3, 4, 5]), 'Respondent ID', 'identity'),
    (pd.Series(['Rahul', 'Rahul', 'Priya', 'Amit']), 'Name', 'identity'),
    (pd.Series(['Rahul', 'Rahul', 'Priya', 'Amit']), 'Student Name', 'identity'),
    (pd.Series([245, 280, 210, 290, 195]), 'Total Marks', 'metadata'),
    (pd.Series([81.5, 93.3, 70.0, 96.6, 65.0]), 'Percentage', 'metadata'),
    (pd.Series([8.5, 9.2, 7.8, 9.6, 6.9]), 'CGPA', 'metadata'),
    (pd.Series([1, 2, 3, 4, 5]), 'Rank', 'metadata'),
    (pd.Series([5, 4, 5, 3, 4]), 'Rating (1-5)', 'metadata'),
    (pd.Series([75, 'AB', 80, 'AB', 85, 'AB']), 'Maths', 'marks'),
    (pd.Series([' 78 ', '45*', '90', '60#', '55']), 'Physics', 'marks'),
    (pd.Series([60, 90, 98, 99, 75]), 'CIT-1', 'marks'),
    (pd.Series([55, 60, 70, 80]), '24CS301.', 'marks'),
]

for s, col_name, expected in test_cases:
    cls, q = classify_column(s, col_name)
    status = "OK" if cls == expected else f"FAIL (got {cls}, expected {expected})"
    print(f"  {col_name:15s} -> {cls:10s} [{status}]")

print("\n--- Testing Reference Pipeline ---")
from services.analyser import analyse_workbook
from services.reporter import generate_report

ref_src = 'Reference/Original_1 ExcelSpreadsheet.xlsx'
ref_json = analyse_workbook(ref_src)
os.makedirs('scratch/test_acc', exist_ok=True)
ref_out = 'scratch/test_acc/ref_report.xlsx'
ref_stats = generate_report(ref_src, ref_json, 'Table 1', ref_out)
print(f"Ref report total_students: {ref_stats['total_students']}, pass: {ref_stats['pass_count']}, fail: {ref_stats['fail_count']}")

print("\n--- Testing cit-1-a_1.xls Pipeline ---")
cit_src = 'uploads/fea5a48088bd4c03ae4f0108ad7d0d54.xlsx'
if os.path.exists(cit_src):
    cit_json = analyse_workbook(cit_src)
    cit_out = 'scratch/test_acc/cit_report.xlsx'
    sheet_name = cit_json[0]['name']
    cit_stats = generate_report(cit_src, cit_json, sheet_name, cit_out)
    print(f"CIT report total_students: {cit_stats['total_students']}, pass: {cit_stats['pass_count']}, fail: {cit_stats['fail_count']}")
    
    # Inspect cit_out rows to check for stray rows
    wb = openpyxl.load_workbook(cit_out, data_only=True)
    ws = wb[sheet_name]
    print(f"CIT output sheet max_row: {ws.max_row}")
    for r in range(48, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        print(f"  Row {r}: {vals}")
